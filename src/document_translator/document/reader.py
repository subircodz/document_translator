"""Read supported structures from DOCX files."""

from pathlib import Path

from docx import Document

from document_translator.document.model import (
    DocumentModel,
    ParagraphModel,
    RunModel,
    TableModel,
)


def _paragraph_model(paragraph) -> ParagraphModel:
    return ParagraphModel(
        text=paragraph.text,
        style=paragraph.style.name if paragraph.style is not None else None,
        runs=tuple(
            RunModel(
                text=run.text,
                bold=run.bold,
                italic=run.italic,
                underline=run.underline,
            )
            for run in paragraph.runs
        ),
    )


def read_docx(path: str | Path) -> DocumentModel:
    """Read paragraphs and tables from a DOCX document."""
    document = Document(path)
    paragraphs = tuple(_paragraph_model(paragraph) for paragraph in document.paragraphs)
    tables = tuple(
        TableModel(
            rows=tuple(
                tuple(
                    _paragraph_model(paragraph)
                    for paragraph in cell.paragraphs
                )
                for row in table.rows
                for cell in row.cells
            )
        )
        for table in document.tables
    )
    return DocumentModel(paragraphs=paragraphs, tables=tables)
