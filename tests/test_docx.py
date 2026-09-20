from docx import Document

from document_translator.document.model import DocumentModel, ParagraphModel
from document_translator.document.reader import read_docx
from document_translator.document.writer import write_docx


def test_docx_round_trip_preserves_order_structure_runs_and_tables(tmp_path) -> None:
    source_path = tmp_path / "source.docx"
    output_path = tmp_path / "output.docx"

    source = Document()
    heading = source.add_heading("English Document", level=1)
    heading.runs[0].bold = True

    paragraph = source.add_paragraph()
    first = paragraph.add_run("Hello ")
    first.bold = True
    second = paragraph.add_run("world")
    second.italic = True

    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Language"
    table.cell(1, 0).text = "Subir"
    table.cell(1, 1).text = "English"

    source.add_paragraph("After table")
    source.save(source_path)

    model = read_docx(source_path)
    write_docx(model, output_path)
    result = read_docx(output_path)

    assert len(result.blocks) == 4
    assert isinstance(result.blocks[0], ParagraphModel)
    assert result.blocks[0].text == "English Document"
    assert result.blocks[0].style == "Heading 1"
    assert result.blocks[1].runs[0].bold is True
    assert result.blocks[1].runs[1].italic is True

    table_model = result.blocks[2]
    assert [[p.text for p in cell] for row in table_model.rows for cell in row] == [
        ["Name"],
        ["Language"],
        ["Subir"],
        ["English"],
    ]
    assert result.blocks[3].text == "After table"


def test_docx_unicode_round_trip(tmp_path) -> None:
    source_path = tmp_path / "unicode.docx"
    output_path = tmp_path / "unicode-output.docx"

    source = Document()
    source.add_paragraph("Hindi नमस्ते")
    source.add_paragraph("Bengali নমস্কার")
    source.add_paragraph("Kannada ಕನ್ನಡ")
    source.add_paragraph("Telugu తెలుగు")
    source.add_paragraph("Tamil தமிழ்")
    source.add_paragraph("Malayalam മലയാളം")
    source.save(source_path)

    write_docx(read_docx(source_path), output_path)
    result = read_docx(output_path)

    assert [block.text for block in result.blocks if isinstance(block, ParagraphModel)] == [
        "Hindi नमस्ते",
        "Bengali নমস্কার",
        "Kannada ಕನ್ನಡ",
        "Telugu తెలుగు",
        "Tamil தமிழ்",
        "Malayalam മലയാളം",
    ]


def test_writer_accepts_domain_model(tmp_path) -> None:
    output_path = tmp_path / "model.docx"
    model = DocumentModel(
        blocks=(ParagraphModel(text="Hello", style="Normal", runs=()),)
    )

    write_docx(model, output_path)

    assert read_docx(output_path).blocks[0].text == "Hello"
