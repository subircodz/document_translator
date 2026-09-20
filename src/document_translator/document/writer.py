"""Write supported document structures to DOCX files."""

from pathlib import Path

from docx import Document

from document_translator.document.model import (
    DocumentModel,
    ParagraphModel,
    RunModel,
    TableModel,
)


def _write_paragraph(document, paragraph_model: ParagraphModel):
    paragraph = document.add_paragraph(style=paragraph_model.style)
    for run_model in paragraph_model.runs:
        run = paragraph.add_run(run_model.text)
        run.bold = run_model.bold
        run.italic = run_model.italic
        run.underline = run_model.underline
    return paragraph


def _write_table(document, table_model: TableModel) -> None:
    row_count = len(table_model.rows)
    column_count = max((len(row) for row in table_model.rows), default=0)
    table = document.add_table(rows=row_count, cols=column_count)

    for row_index, row_model in enumerate(table_model.rows):
        for column_index, paragraph_model in enumerate(row_model):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            for run_model in paragraph_model.runs:
                run = cell.paragraphs[0].add_run(run_model.text)
                run.bold = run_model.bold
                run.italic = run_model.italic
                run.underline = run_model.underline


def write_docx(document_model: DocumentModel, path: str | Path) -> None:
    """Write a DocumentModel to a DOCX file."""
    document = Document()

    for paragraph_model in document_model.paragraphs:
        _write_paragraph(document, paragraph_model)

    for table_model in document_model.tables:
        _write_table(document, table_model)

    document.save(path)
