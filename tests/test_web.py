import io
import time
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from document_translator.models import Language, TranslationResult
from document_translator.web.app import JobStore, TranslationJob, create_app


class FakeProvider:
    def translate_many(self, requests):
        return [
            TranslationResult(
                request.source_language,
                request.target_language,
                request.text,
                f"translated-{request.text}",
            )
            for request in requests
        ]

    def supports(self, source, target):
        return source is Language.ENGLISH and target in {
            Language.HINDI,
            Language.TAMIL,
        }


def fake_factory(_api_key):
    return FakeProvider()


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Hello", level=1)
    document.add_paragraph("Order ORD-ABC-123")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_home_page_lists_targets() -> None:
    response = TestClient(create_app(fake_factory)).get("/")
    assert response.status_code == 200
    assert 'value="hi"' in response.text
    assert 'value="ta"' in response.text


def test_status_missing_job() -> None:
    response = TestClient(create_app(fake_factory)).get("/api/translations/missing")
    assert response.status_code == 404


def test_upload_rejects_non_docx(monkeypatch) -> None:
    monkeypatch.setenv("GOOGLE_TRANSLATE_API_KEY", "test")
    response = TestClient(create_app(fake_factory)).post(
        "/translate",
        files={"file": ("notes.txt", b"hello", "text/plain")},
        data={"targets": "hi"},
    )
    assert response.status_code == 400
    assert "docx" in response.json()["detail"]


def test_upload_rejects_missing_api_key(monkeypatch) -> None:
    monkeypatch.delenv("GOOGLE_TRANSLATE_API_KEY", raising=False)
    response = TestClient(create_app(fake_factory)).post(
        "/translate",
        files={"file": ("notes.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"targets": "hi"},
    )
    assert response.status_code == 503


def test_upload_translates_and_exposes_report(monkeypatch) -> None:
    monkeypatch.setenv("GOOGLE_TRANSLATE_API_KEY", "test")
    client = TestClient(create_app(fake_factory))
    response = client.post(
        "/translate",
        files={"file": ("report.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"targets": ["hi", "ta"]},
    )
    assert response.status_code == 200
    assert "/jobs/" in response.text
    job_id = response.text.split("/jobs/")[1].split('"')[0]
    status = client.get(f"/api/translations/{job_id}")
    assert status.status_code == 200
    payload = status.json()
    assert payload["status"] == "completed"
    assert payload["progress"] == 100
    assert {item["target"] for item in payload["outputs"]} == {"hi", "ta"}
    document_response = client.get(f"/api/translations/{job_id}/files/hi")
    report_response = client.get(f"/api/translations/{job_id}/reports/hi")
    assert document_response.status_code == 200
    assert report_response.status_code == 200
    assert "PASS" in report_response.text


def test_upload_rejects_invalid_docx(monkeypatch) -> None:
    monkeypatch.setenv("GOOGLE_TRANSLATE_API_KEY", "test")
    response = TestClient(create_app(fake_factory)).post(
        "/translate",
        files={"file": ("broken.docx", b"not-a-zip", "application/octet-stream")},
        data={"targets": "hi"},
    )
    assert response.status_code == 400
    assert "valid DOCX" in response.json()["detail"]


def test_expired_job_cleans_up(tmp_path) -> None:
    import tempfile

    temp_dir = tempfile.TemporaryDirectory(prefix="document-translator-test-")
    work_dir = tmp_path / "job"
    work_dir.mkdir()
    marker = work_dir / "marker.txt"
    marker.write_text("x", encoding="utf-8")
    job = TranslationJob(
        "expired",
        "report.docx",
        (Language.HINDI,),
        work_dir,
        temp_dir,
        ttl_seconds=1,
        created_at=time.time() - 2,
    )
    store = JobStore()
    store.add(job)
    assert store.get("expired") is None
    assert not Path(temp_dir.name).exists()
