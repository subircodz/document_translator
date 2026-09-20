import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from document_translator.config import ConfigurationError, Settings
from document_translator.document.reader import read_docx
from document_translator.document.model import DocumentModel, ParagraphModel, RunModel
from document_translator.models import Language, TranslationResult
from document_translator.translation.document import translate_document
from document_translator.web.app import create_app


class FakeProvider:
    def translate_many(self, requests):
        return [
            TranslationResult(
                request.source_language,
                request.target_language,
                request.text,
                f"translated-{request.text}",
            )
            for request in requests
        ]


def fake_factory(_api_key):
    return FakeProvider()


def make_docx() -> bytes:
    document = Document()
    paragraph = document.add_paragraph()
    bold = paragraph.add_run("Hello")
    bold.bold = True
    italic = paragraph.add_run(" world")
    italic.italic = True
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_translation_preserves_run_boundaries_and_formatting() -> None:
    source = DocumentModel(
        blocks=(
            ParagraphModel(
                text="Hello world",
                style="Normal",
                runs=(
                    RunModel("Hello", bold=True),
                    RunModel(" world", italic=True),
                ),
            ),
        )
    )
    translated = translate_document(
        source,
        FakeProvider(),
        Language.HINDI,
    )
    paragraph = translated.blocks[0]
    assert isinstance(paragraph, ParagraphModel)
    assert [run.text for run in paragraph.runs] == [
        "translated-Hello",
        "translated- world",
    ]
    assert paragraph.runs[0].bold is True
    assert paragraph.runs[1].italic is True
    assert paragraph.text == "translated-Hellotranslated- world"


def test_settings_reject_partial_web_credentials(monkeypatch) -> None:
    monkeypatch.setenv("GOOGLE_TRANSLATE_API_KEY", "test")
    monkeypatch.setenv("DOCUMENT_TRANSLATOR_WEB_USERNAME", "admin")
    monkeypatch.delenv("DOCUMENT_TRANSLATOR_WEB_PASSWORD", raising=False)
    with pytest.raises(ConfigurationError):
        Settings.from_environment()


def test_protected_web_app_requires_basic_auth() -> None:
    settings = Settings(
        google_translate_api_key="test",
        web_username="admin",
        web_password="secret",
        rate_limit_requests=20,
    )
    client = TestClient(create_app(fake_factory, settings=settings))
    assert client.get("/health").status_code == 200
    assert client.get("/").status_code == 401
    assert client.get("/", auth=("admin", "secret")).status_code == 200


def test_protected_web_app_translates_with_auth() -> None:
    settings = Settings(
        google_translate_api_key="test",
        web_username="admin",
        web_password="secret",
        rate_limit_requests=20,
    )
    client = TestClient(create_app(fake_factory, settings=settings))
    response = client.post(
        "/translate",
        auth=("admin", "secret"),
        files={
            "file": (
                "report.docx",
                make_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"targets": "hi"},
    )
    assert response.status_code == 200
    job_id = response.text.split("/jobs/")[1].split('"')[0]
    status = client.get(
        f"/api/translations/{job_id}",
        auth=("admin", "secret"),
    )
    assert status.status_code == 200
    assert status.json()["status"] == "completed"
    output = client.get(
        f"/api/translations/{job_id}/files/hi",
        auth=("admin", "secret"),
    )
    assert output.status_code == 200
    model = read_docx(io.BytesIO(output.content))
    paragraph = model.blocks[0]
    assert isinstance(paragraph, ParagraphModel)
    assert paragraph.runs[0].bold is True
    assert paragraph.runs[1].italic is True
